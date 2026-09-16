import csv
import logging
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST

from .models import ETLRun, ETLResult
from mappings.models import Mapping, ColumnMapping, ETLStep
from accounts.decorators import contributor_or_admin_required

logger = logging.getLogger('validations')


@login_required
def etl_list_view(request):
    """List all validation runs."""
    from django.core.paginator import Paginator

    query = request.GET.get('query', '').strip()
    status_filter = request.GET.get('status', '').strip()
    type_filter = request.GET.get('type', '').strip()
    date_filter = request.GET.get('date', '').strip()

    runs_qs = ETLRun.objects.select_related('mapping', 'triggered_by').all()
    if query:
        runs_qs = runs_qs.filter(mapping__name__icontains=query)
    if status_filter:
        runs_qs = runs_qs.filter(status=status_filter)
    if type_filter:
        runs_qs = runs_qs.filter(trigger_type=type_filter)
    if date_filter:
        runs_qs = runs_qs.filter(created_at__date=date_filter)

    paginator = Paginator(runs_qs, 100)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Attach reverse index to each item in current page to display global numbering
    start_index = paginator.count - (page_obj.start_index() - 1)
    for i, run in enumerate(page_obj.object_list):
        run.rev_index = start_index - i

    return render(request, 'validations/list.html', {
        'runs': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'query': query,
        'status_filter': status_filter,
        'type_filter': type_filter,
        'date_filter': date_filter
    })


@login_required
def etl_report_view(request, run_id):
    """View detailed report for a validation run."""
    run = get_object_or_404(
        ETLRun.objects.select_related('mapping', 'triggered_by'),
        id=run_id
    )
    # Calculate dynamic monitor run number matching Monitor list view
    total_count = ETLRun.objects.count()
    runs_after = ETLRun.objects.filter(id__gt=run.id).count()
    run.rev_index = total_count - runs_after

    results = run.results.select_related('column_mapping').all()
    return render(request, 'validations/report.html', {
        'run': run,
        'results': results,
    })


@login_required
def etl_progress_view(request, run_id):
    """View validation progress (for active runs)."""
    run = get_object_or_404(ETLRun, id=run_id)
    if run.status in ('completed', 'failed'):
        return redirect('validations:report', run_id=run.id)
    return render(request, 'validations/progress.html', {'run': run})


# ─── API Endpoints ───────────────────────────────────────────────────────────

@login_required
def api_etl_progress(request, run_id):
    """AJAX: Get validation progress."""
    run = get_object_or_404(ETLRun, id=run_id)
    return JsonResponse({
        'status': run.status,
        'progress': run.progress,
        'total': run.total_checks,
        'passed': run.records_extracted,
        'failed': run.records_loaded,
    })


@login_required
def api_validate_pipeline(request, mapping_id):
    """AJAX endpoint that runs pre-execution validation checks on a pipeline (Mapping)."""
    mapping = get_object_or_404(Mapping, id=mapping_id)
    validations = []
    has_critical_error = False

    from connections.connector import ConnectorEngine
    source_engine = ConnectorEngine(mapping.source_connection)
    target_engine = ConnectorEngine(mapping.target_connection)

    # 1. Source Connection Check
    try:
        success, msg = source_engine.test_connection()
        if success:
            validations.append({'name': 'Source Connection', 'status': 'success', 'message': 'Connection successful.'})
        else:
            validations.append({'name': 'Source Connection', 'status': 'error', 'message': f'Connection failed: {msg}'})
            has_critical_error = True
    except Exception as e:
        validations.append({'name': 'Source Connection', 'status': 'error', 'message': f'Connection error: {str(e)}'})
        has_critical_error = True

    # 2. Target Connection Check
    try:
        success, msg = target_engine.test_connection()
        if success:
            validations.append({'name': 'Target Connection', 'status': 'success', 'message': 'Connection successful.'})
        else:
            validations.append({'name': 'Target Connection', 'status': 'error', 'message': f'Connection failed: {msg}'})
            has_critical_error = True
    except Exception as e:
        validations.append({'name': 'Target Connection', 'status': 'error', 'message': f'Connection error: {str(e)}'})
        has_critical_error = True

    source_cols_meta = []
    target_cols_meta = []

    # 3. Source Table / Custom Query Check
    if mapping.query_type == 'custom_query':
        if mapping.source_connection and mapping.source_connection.is_file:
            validations.append({
                'name': 'Source SQL Query Syntax',
                'status': 'error',
                'message': 'Cannot execute SQL on flat file connections (CSV/Excel/JSON). Please edit this pipeline and switch Query Type to "Table Select".'
            })
            has_critical_error = True
        else:
            query = mapping.custom_query or ''
            if not query.strip():
                validations.append({'name': 'Source Dataset', 'status': 'error', 'message': 'Custom query is empty.'})
                has_critical_error = True
            else:
                # Check read-only SQL safety
                import re
                clean_query = re.sub(r'--.*$', '', query, flags=re.MULTILINE)
                clean_query = re.sub(r'/\*.*?\*/', '', clean_query, flags=re.DOTALL)
                forbidden_pattern = re.compile(r'\b(insert|update|delete|drop|alter|truncate)\b', re.IGNORECASE)
                match = forbidden_pattern.search(clean_query)
                if match:
                    validations.append({'name': 'Source Query Safety', 'status': 'error', 'message': f"Destructive SQL operation detected in custom query: '{match.group(1).upper()}'."})
                    has_critical_error = True
                else:
                    # Syntax Check: Try to execute query with LIMIT 0
                    try:
                        if not source_engine.is_mocked():
                            db_type = str(mapping.source_connection.connection_type).lower()
                            if db_type == 'oracle':
                                check_query = f"SELECT * FROM ({query}) WHERE ROWNUM = 0"
                            elif db_type == 'db2':
                                check_query = f"SELECT * FROM ({query}) AS temp FETCH FIRST 0 ROWS ONLY"
                            elif db_type in ('mssql', 'sqlserver'):
                                check_query = f"SELECT TOP 0 * FROM ({query}) AS temp"
                            else:
                                check_query = f"SELECT * FROM ({query}) LIMIT 0"
                            source_engine.execute_query(check_query)
                        validations.append({'name': 'Source SQL Query Syntax', 'status': 'success', 'message': 'Custom SQL query syntax is valid.'})
                    except Exception as sqle:
                        validations.append({'name': 'Source SQL Query Syntax', 'status': 'error', 'message': f'SQL query validation failed: {str(sqle)}'})
                        has_critical_error = True
    else:
        # Table Select mode
        if not mapping.source_table:
            validations.append({'name': 'Source Table', 'status': 'error', 'message': 'Source table is not selected.'})
            has_critical_error = True
        else:
            try:
                if not source_engine.is_mocked():
                    tables = source_engine.get_tables(schema=mapping.source_schema or None, catalog=mapping.source_catalog or None)
                    if mapping.source_table not in tables and mapping.source_table.upper() not in [t.upper() for t in tables]:
                        validations.append({'name': 'Source Table Existence', 'status': 'error', 'message': f"Table '{mapping.source_table}' was not found in schema/catalog."})
                        has_critical_error = True
                    else:
                        validations.append({'name': 'Source Table Existence', 'status': 'success', 'message': f"Table '{mapping.source_table}' exists."})
                        source_cols_meta = source_engine.get_columns(schema=mapping.source_schema or None, table=mapping.source_table, catalog=mapping.source_catalog or None)
                else:
                    source_cols_meta = source_engine.get_columns(table=mapping.source_table)
                    validations.append({'name': 'Source Table Existence', 'status': 'success', 'message': f"Table '{mapping.source_table}' verified (mocked connection)."})
            except Exception as e:
                validations.append({'name': 'Source Table Check', 'status': 'warning', 'message': f'Could not verify source table: {str(e)}'})

    # 4. Target Table Check
    if not mapping.target_table:
        validations.append({'name': 'Target Table', 'status': 'error', 'message': 'Target table is not selected.'})
        has_critical_error = True
    else:
        try:
            if not target_engine.is_mocked():
                tables = target_engine.get_tables(schema=mapping.target_schema or None, catalog=mapping.target_catalog or None)
                if mapping.target_table not in tables and mapping.target_table.upper() not in [t.upper() for t in tables]:
                    validations.append({'name': 'Target Table Existence', 'status': 'error', 'message': f"Table '{mapping.target_table}' was not found in target schema."})
                    has_critical_error = True
                else:
                    validations.append({'name': 'Target Table Existence', 'status': 'success', 'message': f"Table '{mapping.target_table}' exists."})
                    target_cols_meta = target_engine.get_columns(schema=mapping.target_schema or None, table=mapping.target_table, catalog=mapping.target_catalog or None)
            else:
                target_cols_meta = target_engine.get_columns(table=mapping.target_table)
                validations.append({'name': 'Target Table Existence', 'status': 'success', 'message': f"Table '{mapping.target_table}' verified (mocked connection)."})
        except Exception as e:
            validations.append({'name': 'Target Table Check', 'status': 'warning', 'message': f'Could not verify target table: {str(e)}'})

    # 5. Mappings Check
    mappings_count = mapping.column_mappings.count()
    if mappings_count == 0:
        validations.append({'name': 'Column Mappings', 'status': 'error', 'message': 'No columns mapped between source and target.'})
        has_critical_error = True
    else:
        validations.append({'name': 'Column Mappings', 'status': 'success', 'message': f'{mappings_count} columns mapped.'})

        # 6. Source and Target Column Existence & Datatype Compatibility Check
        src_cols_dict = {c['name'].lower(): c['type'] for c in source_cols_meta} if source_cols_meta else {}
        tgt_cols_dict = {c['name'].lower(): c['type'] for c in target_cols_meta} if target_cols_meta else {}

        column_existence_passed = True
        datatype_issues = []

        for cm in mapping.column_mappings.all():
            # Check Source column
            if source_cols_meta and mapping.query_type != 'custom_query':
                if cm.source_column.lower() not in src_cols_dict:
                    validations.append({'name': 'Column Existence', 'status': 'error', 'message': f"Source column '{cm.source_column}' does not exist in table."})
                    has_critical_error = True
                    column_existence_passed = False
            
            # Check Target column
            if target_cols_meta:
                if cm.target_column.lower() not in tgt_cols_dict:
                    validations.append({'name': 'Column Existence', 'status': 'error', 'message': f"Target column '{cm.target_column}' does not exist in table."})
                    has_critical_error = True
                    column_existence_passed = False

            # Datatype compatibility
            if source_cols_meta and target_cols_meta and mapping.query_type != 'custom_query':
                src_type = src_cols_dict.get(cm.source_column.lower(), '')
                tgt_type = tgt_cols_dict.get(cm.target_column.lower(), '')
                if src_type and tgt_type:
                    src_cat = get_datatype_category(src_type)
                    tgt_cat = get_datatype_category(tgt_type)
                    if src_cat != tgt_cat:
                        datatype_issues.append(f"Mapped '{cm.source_column}' ({src_type}) to '{cm.target_column}' ({tgt_type}).")

        if column_existence_passed and mappings_count > 0:
            validations.append({'name': 'Column Existence', 'status': 'success', 'message': 'All mapped columns exist in source and target.'})

        if datatype_issues:
            msg = "Datatype mismatch detected (may cause conversion failures):<br>" + "<br>".join(datatype_issues)
            validations.append({'name': 'Datatype Compatibility', 'status': 'warning', 'message': msg})
        else:
            if mappings_count > 0:
                validations.append({'name': 'Datatype Compatibility', 'status': 'success', 'message': 'All mapped columns have compatible datatypes.'})

    # 7. Table Filter check
    if mapping.filter_column and mapping.query_type != 'custom_query':
        if source_cols_meta and mapping.filter_column.lower() not in src_cols_dict:
            validations.append({'name': 'Filter Column Check', 'status': 'error', 'message': f"Filter column '{mapping.filter_column}' does not exist in source table."})
            has_critical_error = True
        else:
            validations.append({'name': 'Filter Column Check', 'status': 'success', 'message': f"Filter column '{mapping.filter_column}' verified."})

    # 8. SCD Configuration Validation
    if mapping.load_mode in ('scd1', 'scd2'):
        if not mapping.scd_business_key:
            validations.append({'name': 'SCD Business Key', 'status': 'error', 'message': 'Business key is not specified for SCD strategy.'})
            has_critical_error = True
        else:
            if target_cols_meta:
                if mapping.scd_business_key.lower() not in tgt_cols_dict:
                    validations.append({'name': 'SCD Business Key Existence', 'status': 'error', 'message': f"Business key '{mapping.scd_business_key}' does not exist in target table."})
                    has_critical_error = True
                else:
                    validations.append({'name': 'SCD Business Key', 'status': 'success', 'message': f"Business key '{mapping.scd_business_key}' exists in target."})

        # Columns to track
        track_cols = [c.strip() for c in (mapping.scd_track_columns or '').split(',') if c.strip()]
        if not track_cols:
            validations.append({'name': 'SCD Tracked Columns', 'status': 'error', 'message': 'No columns specified to track/update for SCD.'})
            has_critical_error = True
        else:
            invalid_track_cols = []
            if target_cols_meta:
                for col in track_cols:
                    if col.lower() not in tgt_cols_dict:
                        invalid_track_cols.append(col)
            if invalid_track_cols:
                validations.append({'name': 'SCD Tracked Columns Existence', 'status': 'error', 'message': f"Tracked columns do not exist in target: {', '.join(invalid_track_cols)}"})
                has_critical_error = True
            else:
                validations.append({'name': 'SCD Tracked Columns', 'status': 'success', 'message': f'{len(track_cols)} tracked columns verified.'})

        # SCD 2 specific columns
        if mapping.load_mode == 'scd2':
            for col_name, field_label in [
                (mapping.scd_effective_from, 'Effective From'),
                (mapping.scd_effective_to, 'Effective To'),
                (mapping.scd_active_flag, 'Current/Active Flag')
            ]:
                if not col_name:
                    validations.append({'name': f'SCD {field_label}', 'status': 'error', 'message': f'{field_label} column is not specified.'})
                    has_critical_error = True
                elif target_cols_meta:
                    if col_name.lower() not in tgt_cols_dict:
                        validations.append({'name': f'SCD {field_label} Existence', 'status': 'error', 'message': f"{field_label} column '{col_name}' does not exist in target table."})
                        has_critical_error = True
                    else:
                        validations.append({'name': f'SCD {field_label}', 'status': 'success', 'message': f"{field_label} column '{col_name}' exists in target."})

    return JsonResponse({
        'success': not has_critical_error,
        'validations': validations
    })


@login_required
@contributor_or_admin_required
@require_POST
def api_trigger_etl(request, mapping_id):
    """AJAX: Manually trigger a validation for a mapping."""
    if hasattr(request.user, 'profile') and request.user.profile.role == 'auditor':
        return JsonResponse({'success': False, 'error': 'Permission denied: Auditor cannot trigger validations.'}, status=403)
    mapping = get_object_or_404(Mapping, id=mapping_id)

    # Extract JSON parameters if present
    import json
    parameters = {}
    if request.content_type == 'application/json':
        try:
            body = json.loads(request.body)
            parameters = body.get('parameters', {})
        except Exception:
            pass
    else:
        param_str = request.POST.get('parameters', '{}')
        try:
            parameters = json.loads(param_str)
        except Exception:
            parameters = {}

    run = ETLRun.objects.create(
        mapping=mapping,
        triggered_by=request.user,
        trigger_type='manual',
        status='pending',
        parameters=parameters,
    )

    # Try to use Celery, fall back to synchronous
    try:
        from .tasks import run_etl_task
        run_etl_task.delay(run.id)
    except Exception:
        # Fallback: run synchronously
        from .engine import ETLEngine
        engine = ETLEngine(run)
        try:
            engine.execute()
            if run.triggered_by:
                from dashboard.models import Notification
                status_text = "Passed" if run.failed_checks == 0 else "Failed"
                Notification.objects.create(
                    user=run.triggered_by,
                    title=f"Validation Run {run.id} Completed",
                    message=f"Pipeline: {run.mapping.name}\nStatus: {status_text} ({run.passed_checks}/{run.total_checks} checks passed)",
                    level='success' if run.failed_checks == 0 else 'warning'
                )
        except Exception as e:
            logger.error(f"Sync validation failed: {e}")
            if run.triggered_by:
                from dashboard.models import Notification
                Notification.objects.create(
                    user=run.triggered_by,
                    title=f"Validation Run {run.id} Failed",
                    message=f"Pipeline: {run.mapping.name}\nError: {e}",
                    level='error'
                )

    return JsonResponse({'success': True, 'run_id': run.id})


@login_required
def export_report(request, run_id):
    """Export validation report as Excel file."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    
    run = get_object_or_404(ETLRun, id=run_id)
    mapping = run.mapping
    results = run.results.select_related('column_mapping').all()

    src_conn = mapping.source_connection
    col_headers = [
        'Table Name',
        'Column Name',
        'Source Validation Operation(operation choosen)',
        'Target Validation Operation (operation choosen)',
        'Result',
        'Difference'
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validation Report"

    # Write headers
    ws.append(col_headers)

    # Style header row
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Deep Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        res_val = "MATCH" if r.is_match else "MISMATCH"
        col_name = r.source_column or (r.column_mapping.source_column if r.column_mapping else 'Unknown')
        ws.append([
            mapping.source_table.upper(),
            col_name,
            r.source_op_display,
            r.target_op_display,
            res_val,
            r.difference
        ])

    # Formatting columns
    for row in range(2, ws.max_row + 1):
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            if col in [1, 2, 3, 4]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            if col == 5:
                if cell.value == "MATCH":
                    cell.font = Font(name="Calibri", size=11, bold=True, color="15803D") # Green
                else:
                    cell.font = Font(name="Calibri", size=11, bold=True, color="B91C1C") # Red

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    val_date = run.created_at.strftime('%Y-%m-%d')
    import re
    safe_pipeline_name = re.sub(r'[\\/*?:"<>|]', "", mapping.name).strip()
    filename = f"{safe_pipeline_name}.xlsx"

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)

    # Log report generation audit
    try:
        from logs.models import AuditLog
        AuditLog.objects.create(
            user=request.user,
            action=f"Report Generated: {mapping.source_table} on {val_date}",
            entity_type="ETLRun",
            entity_id=run.id,
            level="success",
        )
    except Exception:
        pass

    return response


def get_datatype_category(type_str, name_str=''):
    t = str(type_str or '').upper()
    n = str(name_str or '').upper()
    if any(x in t for x in ('INT', 'BIGINT', 'SMALLINT', 'TINYINT', 'NUMERIC', 'DECIMAL', 'FLOAT', 'DOUBLE', 'REAL', 'NUMBER')):
        return 'INTEGER'
    elif (any(x in t for x in ('DATE', 'TIME', 'TIMESTAMP')) or 
          any(x in n for x in ('DATE', 'TIME', 'TIMESTAMP')) or 
          n.endswith('_AT') or n.endswith('_ON') or n == 'AT' or 'DT' in n):
        return 'DATE'
    elif any(x in t for x in ('BOOL', 'BOOLEAN')):
        return 'BOOLEAN'
    else:
        return 'VARCHAR'

def get_applicable_operations(category):
    if category == 'INTEGER':
        return ['null_check', 'sum', 'avg', 'min', 'max', 'range_check', 'duplicate_check', 'count', 'row_count', 'unique_check', 'distinct_count', 'data_type_check', 'std_dev', 'variance', 'median', 'mode']
    elif category == 'DATE':
        return ['min_date', 'max_date', 'null_check', 'duplicate_check', 'count', 'row_count', 'unique_check', 'distinct_count']
    elif category == 'BOOLEAN':
        return ['null_check', 'count', 'row_count', 'duplicate_check', 'unique_check', 'distinct_count']
    else: # VARCHAR
        return ['null_check', 'length_sum_check', 'sum_length', 'regex_check', 'duplicate_check', 'unique_check', 'distinct_count', 'row_count', 'count', 'pattern_match', 'data_type_check']

@login_required
@contributor_or_admin_required
def quick_etl_view(request):
    """Create a quick mapping and trigger validation from the dashboard."""
    if request.method == 'POST':
        try:
            source_conn_id = request.POST.get('source_connection')
            source_catalog = request.POST.get('source_catalog', '')
            source_schema = request.POST.get('source_schema', '')
            source_table = request.POST.get('source_table', '')
            target_conn_id = request.POST.get('target_connection')
            target_catalog = request.POST.get('target_catalog', '')
            target_schema = request.POST.get('target_schema', '')
            target_table = request.POST.get('target_table', '')
            load_mode = request.POST.get('load_mode', 'truncate')
            try:
                batch_size = int(request.POST.get('batch_size', 10000))
            except (ValueError, TypeError):
                batch_size = 10000
            
            query_type = request.POST.get('query_type', 'table')
            custom_query = request.POST.get('custom_query', '')
            if query_type == 'custom_query':
                source_table = 'Custom Query'

            # Resolving Source Date Filters
            source_date_column = request.POST.get('source_date_column', '')
            source_date_filter_type = request.POST.get('source_date_filter_type', 'none')
            if not source_date_column:
                source_date_filter_type = 'none'
            source_date_filter_start = None
            source_date_filter_end = None
            source_date_value_type = request.POST.get('source_date_value_type', 'calendar')
            source_date_relative_operator = request.POST.get('source_date_relative_operator', '+')
            try:
                source_date_relative_value = int(request.POST.get('source_date_relative_value', 0) or 0)
            except (ValueError, TypeError):
                source_date_relative_value = 0
            source_date_operator = request.POST.get('source_date_operator', '=')

            if source_date_filter_type == 'specific':
                if source_date_value_type == 'calendar':
                    source_date_single = request.POST.get('source_date_single')
                    if source_date_single:
                        source_date_filter_start = source_date_single
                        source_date_filter_end = source_date_single
            elif source_date_filter_type == 'range':
                source_date_filter_start = request.POST.get('source_date_filter_start') or None
                source_date_filter_end = request.POST.get('source_date_filter_end') or None

            # Resolving Target Date Filters
            target_date_column = request.POST.get('target_date_column', '')
            target_date_filter_type = request.POST.get('target_date_filter_type', 'none')
            if not target_date_column:
                target_date_filter_type = 'none'
            target_date_filter_start = None
            target_date_filter_end = None
            target_date_value_type = request.POST.get('target_date_value_type', 'calendar')
            target_date_relative_operator = request.POST.get('target_date_relative_operator', '+')
            try:
                target_date_relative_value = int(request.POST.get('target_date_relative_value', 0) or 0)
            except (ValueError, TypeError):
                target_date_relative_value = 0
            target_date_operator = request.POST.get('target_date_operator', '=')

            if target_date_filter_type == 'specific':
                if target_date_value_type == 'calendar':
                    target_date_single = request.POST.get('target_date_single')
                    if target_date_single:
                        target_date_filter_start = target_date_single
                        target_date_filter_end = target_date_single
            elif target_date_filter_type == 'range':
                target_date_filter_start = request.POST.get('target_date_filter_start') or None
                target_date_filter_end = request.POST.get('target_date_filter_end') or None

            # Create a quick mapping
            from django.utils import timezone
            now_str = timezone.now().strftime('%Y-%m-%d %H:%M')
            if query_type == 'custom_query':
                quick_name = f"Quick Validate: Custom Query -> {target_table} ({now_str})".strip()
            else:
                quick_name = f"Quick Validate: {source_table} -> {target_table} ({now_str})".strip()
            
            mapping_data = {
                'name': quick_name,
                'description': "Triggered from Dashboard Quick Workspace",
                'source_connection_id': source_conn_id,
                'source_catalog': source_catalog,
                'source_schema': source_schema,
                'source_table': source_table,
                'target_connection_id': target_conn_id,
                'target_catalog': target_catalog,
                'target_schema': target_schema,
                'target_table': target_table,
                'load_mode': load_mode,
                'batch_size': batch_size,
                'created_by': request.user,
                'query_type': query_type,
                'custom_query': custom_query,
                'filter_column': request.POST.get('filter_column', ''),
                'filter_condition': request.POST.get('filter_condition', ''),
                'incremental_column': request.POST.get('incremental_column', ''),
                'incremental_value': request.POST.get('incremental_value', ''),
                'scd_business_key': request.POST.get('scd_business_key', ''),
                'scd_track_columns': ','.join(request.POST.getlist('scd_track_columns')),
                'scd_effective_from': request.POST.get('scd_effective_from', ''),
                'scd_effective_to': request.POST.get('scd_effective_to', ''),
                'scd_active_flag': request.POST.get('scd_active_flag', ''),
                'pre_sql': request.POST.get('pre_sql', ''),
                'pre_sql_location': request.POST.get('pre_sql_location', 'target'),
                'post_sql': request.POST.get('post_sql', ''),
                'post_sql_location': request.POST.get('post_sql_location', 'target'),
                'source_date_column': source_date_column,
                'source_date_filter_type': source_date_filter_type,
                'source_date_filter_start': source_date_filter_start,
                'source_date_filter_end': source_date_filter_end,
                'source_date_operator': source_date_operator,
                'target_date_column': target_date_column,
                'target_date_filter_type': target_date_filter_type,
                'target_date_filter_start': target_date_filter_start,
                'target_date_filter_end': target_date_filter_end,
                'target_date_operator': target_date_operator,
                'source_date_range_operator_start': request.POST.get('source_date_range_operator_start', '>='),
                'source_date_range_operator_end': request.POST.get('source_date_range_operator_end', '<='),
                'target_date_range_operator_start': request.POST.get('target_date_range_operator_start', '>='),
                'target_date_range_operator_end': request.POST.get('target_date_range_operator_end', '<='),
                # Stale fields that might be submitted from client or UI
                'source_date_value_type': source_date_value_type,
                'source_date_relative_operator': source_date_relative_operator,
                'source_date_relative_value': source_date_relative_value,
                'target_date_value_type': target_date_value_type,
                'target_date_relative_operator': target_date_relative_operator,
                'target_date_relative_value': target_date_relative_value,
            }

            # Dynamic fields verification and defensive logging
            model_fields = set()
            for f in Mapping._meta.get_fields():
                model_fields.add(f.name)
                if hasattr(f, 'attname'):
                    model_fields.add(f.attname)

            rejected_fields = {}
            for key in list(mapping_data.keys()):
                if key not in model_fields:
                    rejected_fields[key] = mapping_data[key]
                    logger.warning(
                        f"Field mismatch: Submitted field '{key}' is not a valid attribute of the Mapping model. "
                        f"Removing from parameters list. "
                        f"Model fields: {sorted(list(model_fields))}"
                    )
                    del mapping_data[key]

            try:
                mapping = Mapping.objects.create(**mapping_data)
            except Exception as e:
                logger.error(
                    f"Quick validation creation error: {e}. "
                    f"Submitted keys: {list(mapping_data.keys())}. "
                    f"Rejected data: {rejected_fields}. "
                    f"Model fields: {sorted(list(model_fields))}"
                )
                raise

            try:
                from dashboard.models import FormDraft
                FormDraft.objects.filter(user=request.user, page_key='quick_validate', status='draft').update(status='completed')
            except Exception:
                pass

            # Read columns json
            column_data = request.POST.get('column_mappings_json', '[]')
            import json
            try:
                columns = json.loads(column_data)
            except json.JSONDecodeError:
                columns = []

            # Fallback values if JSON is empty but old format is present
            source_cols = request.POST.getlist('source_columns')
            target_cols = request.POST.getlist('target_columns')
            selected_ops = request.POST.getlist('operations')

            # If "__all__" is passed in columns mapping
            if (columns and columns[0].get('source_column') == '__all__') or ('__all__' in source_cols or '__all__' in target_cols):
                user_selected_ops = []
                if columns and columns[0].get('source_column') == '__all__':
                    user_selected_ops = columns[0].get('operations', [])
                elif selected_ops:
                    user_selected_ops = selected_ops

                from connections.connector import ConnectorEngine
                source_conn = mapping.source_connection
                target_conn = mapping.target_connection
                source_engine = ConnectorEngine(source_conn)
                target_engine = ConnectorEngine(target_conn)
                
                src_all_cols = source_engine.get_columns(source_schema if source_schema != 'file' else None, source_table, catalog=source_catalog or None)
                tgt_all_cols = target_engine.get_columns(target_schema if target_schema != 'file' else None, target_table, catalog=target_catalog or None)
                
                expanded_columns = []
                for s_col in src_all_cols:
                    matched_t = next((t_col for t_col in tgt_all_cols if t_col['name'].lower() == s_col['name'].lower()), None)
                    if matched_t:
                        s_cat = get_datatype_category(s_col['type'], s_col['name'])
                        all_ops = get_applicable_operations(s_cat)
                        if user_selected_ops:
                            ops = [op for op in all_ops if op in user_selected_ops]
                        else:
                            ops = all_ops
                        expanded_columns.append({
                            'source_column': s_col['name'],
                            'source_datatype': s_col['type'],
                            'target_column': matched_t['name'],
                            'target_datatype': matched_t['type'],
                            'operations': ops
                        })
                columns = expanded_columns
            
            if not columns and source_cols and target_cols:
                # Old manual pairing fallback
                max_len = max(len(source_cols), len(target_cols))
                for i in range(max_len):
                    src = source_cols[i] if i < len(source_cols) else ''
                    tgt = target_cols[i] if i < len(target_cols) else ''
                    if src and tgt:
                        columns.append({
                            'source_column': src,
                            'source_datatype': 'unknown',
                            'target_column': tgt,
                            'target_datatype': 'unknown',
                            'operations': selected_ops
                        })

            # Create column mappings and rules
            for col in columns:
                s_col = col.get('source_column', '')
                t_col = col.get('target_column', '')
                s_type = col.get('source_datatype', 'unknown')
                t_type = col.get('target_datatype', 'unknown')
                
                col_mapping = ColumnMapping.objects.create(
                    mapping=mapping,
                    source_column=s_col,
                    source_datatype=s_type,
                    target_column=t_col,
                    target_datatype=t_type,
                )
                
                for op in col.get('operations', []):
                    ETLStep.objects.create(
                        column_mapping=col_mapping,
                        operation=op,
                    )
            
            # Create Validation Run
            param_str = request.POST.get('parameters', '{}')
            try:
                parameters = json.loads(param_str)
            except Exception:
                parameters = {}

            run = ETLRun.objects.create(
                mapping=mapping,
                triggered_by=request.user,
                trigger_type='manual',
                status='pending',
                source_date_filter_start=source_date_filter_start,
                source_date_filter_end=source_date_filter_end,
                target_date_filter_start=target_date_filter_start,
                target_date_filter_end=target_date_filter_end,
                parameters=parameters,
            )
            
            # Execute validation run (eager or async)
            try:
                from .tasks import run_etl_task
                run_etl_task.delay(run.id)
            except Exception:
                from .engine import ETLEngine
                engine = ETLEngine(run)
                try:
                    engine.execute()
                    if run.triggered_by:
                        from dashboard.models import Notification
                        status_text = "Success"
                        Notification.objects.create(
                            user=run.triggered_by,
                            title=f"ETL Run {run.id} Completed",
                            message=f"Pipeline: {run.mapping.name}\nStatus: {status_text} (Extracted: {run.records_extracted} / Loaded: {run.records_loaded} rows)",
                            level='success'
                        )
                except Exception as e:
                    logger.error(f"Sync quick validation failed: {e}")
                    if run.triggered_by:
                        from dashboard.models import Notification
                        Notification.objects.create(
                            user=run.triggered_by,
                            title=f"ETL Run {run.id} Failed",
                            message=f"Pipeline: {run.mapping.name}\nError: {e}",
                            level='error'
                        )
            
            messages.success(request, 'Quick ETL pipeline triggered successfully!')
            return redirect('validations:progress', run_id=run.id)
            
        except Exception as e:
            logger.error(f"Quick ETL execution error: {e}")
            messages.error(request, f"Failed to run quick ETL: {str(e)}")
            return redirect('dashboard:index')
            
    return redirect('dashboard:index')


@login_required
def api_mapping_rules_metadata(request, mapping_id):
    """AJAX: Get validation rules requiring user parameters."""
    mapping = get_object_or_404(Mapping, id=mapping_id)
    rules_needing_params = []
    
    column_mappings = mapping.column_mappings.all()
    for cm in column_mappings:
        for rule in cm.rules.filter(is_active=True):
            if rule.operation in ('contains_check', 'pattern_match'):
                rules_needing_params.append({
                    'id': rule.id,
                    'column': cm.source_column,
                    'operation': rule.operation,
                    'operation_display': rule.get_operation_display(),
                })
                
    return JsonResponse({
        'requires_parameters': len(rules_needing_params) > 0,
        'rules': rules_needing_params
    })


@login_required
@contributor_or_admin_required
def etl_delete_view(request, run_id):
    """Delete a validation run."""
    run = get_object_or_404(ETLRun, id=run_id)
    if request.method == 'POST':
        run.delete()
        messages.success(request, f'Validation Run {run_id} deleted successfully.')
    
    referer = request.META.get('HTTP_REFERER')
    if referer and 'report' not in referer and 'progress' not in referer:
        return redirect(referer)
    return redirect('validations:list')


@login_required
def pipeline_monitor_history_view(request, mapping_id):
    """View monitor run history for a specific pipeline."""
    mapping = get_object_or_404(
        Mapping.objects.select_related('source_connection', 'target_connection', 'created_by'),
        id=mapping_id
    )
    runs = list(mapping.etl_runs.select_related('triggered_by').all().order_by('-created_at'))
    total_count = ETLRun.objects.count()
    for run in runs:
        runs_after = ETLRun.objects.filter(id__gt=run.id).count()
        run.rev_index = total_count - runs_after

    return render(request, 'validations/pipeline_history.html', {
        'mapping': mapping,
        'runs': runs,
    })


@login_required
@contributor_or_admin_required
@require_POST
def api_send_report_email(request, run_id):
    """AJAX: Send validation report to a manually specified email."""
    if hasattr(request.user, 'profile') and request.user.profile.role == 'auditor':
        return JsonResponse({'success': False, 'error': 'Permission denied: Auditors cannot trigger emails.'}, status=403)
        
    import json
    email = None
    if request.content_type == 'application/json':
        try:
            body = json.loads(request.body)
            email = body.get('email', '').strip()
        except Exception:
            pass
    else:
        email = request.POST.get('email', '').strip()

    if not email:
        return JsonResponse({'success': False, 'error': 'Email address is required.'}, status=400)

    if '@' not in email:
        return JsonResponse({'success': False, 'error': 'Invalid email address format.'}, status=400)

    run = get_object_or_404(ETLRun, id=run_id)
    
    try:
        from notifications.email_service import send_validation_email
        notification = send_validation_email(run, recipient_email=email)
        if notification and notification.sent_status == 'success':
            return JsonResponse({'success': True, 'message': f'Report email successfully sent to {email}'})
        else:
            err = 'Unknown error'
            if notification:
                err = notification.error_message
                is_none_err = (
                    not err or 
                    str(err).strip().lower() in ('none', '', '(none, none)', 'none, none', '(none,)', 'none,', '("none", "none")') or
                    ('none' in str(err).strip().lower() and ('(' in str(err) or ',' in str(err)))
                )
                if is_none_err:
                    err = "SMTP/Connection Error: Connection to the mail server failed. Please check your SMTP settings in settings.py / .env and VDI network permissions."
            return JsonResponse({'success': False, 'error': f'Failed to send email: {err}'})
    except Exception as e:
        logger.error(f"Error sending manual email: {e}")
        return JsonResponse({'success': False, 'error': f'Server error: {str(e)}'}, status=500)


